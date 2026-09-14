"""Extract plain text from a judgment file (.pdf, .docx, .txt, .md)."""

from __future__ import annotations

from pathlib import Path


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise RuntimeError(
            "PDF 파일을 읽으려면 pypdf가 필요합니다. "
            "다음 명령으로 설치하세요:  pip install pypdf"
        ) from e
    reader = PdfReader(str(path))
    parts: list[str] = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        parts.append(f"[Page {i}]\n{text}")
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError(
            "DOCX 파일을 읽으려면 python-docx가 필요합니다. "
            "다음 명령으로 설치하세요:  pip install python-docx"
        ) from e
    doc = Document(str(path))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def read(path_str: str) -> str:
    path = Path(path_str)
    if not path.exists():
        raise FileNotFoundError(f"파일을 찾을 수 없습니다: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix in {".docx", ".dotx"}:
        return _read_docx(path)
    if suffix in {".txt", ".md", ".text"}:
        return path.read_text(encoding="utf-8", errors="replace")

    raise ValueError(
        f"지원하지 않는 파일 형식: {suffix}. "
        "지원 형식: .pdf, .docx, .txt, .md"
    )
